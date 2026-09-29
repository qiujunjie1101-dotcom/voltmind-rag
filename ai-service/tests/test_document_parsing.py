"""四种文档解析器的结构、异常和资源限制测试。"""

from __future__ import annotations

import pytest
from docx.document import Document as DocxDocument

from app.exceptions import (
    CorruptDocumentError,
    DocumentTooLargeError,
    EmptyDocumentError,
    EncryptedPdfError,
    NoTextContentError,
    ParsingResourceLimitError,
    TextDecodingError,
    UnsupportedDocumentTypeError,
)
from app.parsing import DocumentParser, ParsingLimits
from tests.parsing_samples import docx_bytes, encrypted_pdf_bytes, pdf_bytes


@pytest.fixture(scope="module")
def parser() -> DocumentParser:
    return DocumentParser()


def test_parses_text_pdf(parser: DocumentParser) -> None:
    result = parser.parse(11, "manual.pdf", pdf_bytes("First PDF page"))

    assert result.document_id == 11
    assert result.file_type == "PDF"
    assert len(result.blocks) == 1
    assert "First PDF page" in result.blocks[0].text
    assert result.blocks[0].page_number == 1
    assert result.blocks[0].section is None


def test_preserves_pdf_page_order_and_real_page_numbers(parser: DocumentParser) -> None:
    result = parser.parse(
        12,
        "multi.pdf",
        pdf_bytes("Page one", None, "Page three"),
    )

    assert [block.page_number for block in result.blocks] == [1, 3]
    assert "Page one" in result.blocks[0].text
    assert "Page three" in result.blocks[1].text


def test_pdf_without_extractable_text_is_rejected(parser: DocumentParser) -> None:
    with pytest.raises(NoTextContentError):
        parser.parse(13, "blank.pdf", pdf_bytes(None))


@pytest.mark.parametrize(
    ("content", "expected"),
    [
        pytest.param(b"%PDF-1.7\nnot a real pdf", CorruptDocumentError, id="corrupt"),
        pytest.param(encrypted_pdf_bytes(), EncryptedPdfError, id="encrypted"),
    ],
)
def test_rejects_corrupt_or_encrypted_pdf(
    parser: DocumentParser,
    content: bytes,
    expected: type[Exception],
) -> None:
    with pytest.raises(expected):
        parser.parse(14, "unsafe.pdf", content)


def test_pdf_page_limit_stops_abnormal_document() -> None:
    parser = DocumentParser(limits=ParsingLimits(max_pdf_pages=1))

    with pytest.raises(ParsingResourceLimitError):
        parser.parse(15, "large.pdf", pdf_bytes("one", "two"))


def test_parses_docx_paragraphs(parser: DocumentParser) -> None:
    content = docx_bytes(
        lambda document: (
            document.add_paragraph("First paragraph"),
            document.add_paragraph("Second paragraph"),
        )
    )

    result = parser.parse(21, "notes.docx", content)

    assert result.file_type == "DOCX"
    assert [block.text for block in result.blocks] == ["First paragraph", "Second paragraph"]
    assert all(block.page_number is None for block in result.blocks)


def test_docx_preserves_heading_paragraph_table_order(parser: DocumentParser) -> None:
    def build(document: DocxDocument) -> None:
        document.add_heading("Overview", level=1)
        document.add_paragraph("Before table")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "alpha"
        table.cell(1, 1).text = "1"
        document.add_heading("Details", level=2)
        document.add_paragraph("After table")

    result = parser.parse(22, "structured.docx", docx_bytes(build))

    assert [block.text for block in result.blocks] == [
        "Overview",
        "Before table",
        "Name | Value\nalpha | 1",
        "Details",
        "After table",
    ]
    assert [block.section for block in result.blocks] == [
        "Overview",
        "Overview",
        "Overview",
        "Overview > Details",
        "Overview > Details",
    ]


def test_rejects_corrupt_docx(parser: DocumentParser) -> None:
    with pytest.raises(CorruptDocumentError):
        parser.parse(23, "broken.docx", b"not a zip archive")


def test_docx_uncompressed_size_is_limited() -> None:
    parser = DocumentParser(limits=ParsingLimits(max_docx_uncompressed_bytes=1))
    content = docx_bytes(lambda document: document.add_paragraph("text"))

    with pytest.raises(ParsingResourceLimitError):
        parser.parse(24, "large.docx", content)


def test_markdown_preserves_heading_hierarchy_and_code_block(parser: DocumentParser) -> None:
    content = b"""# Guide
Intro paragraph.

## Example
```python
def hello():
    return "ok"
```

After code.
"""

    result = parser.parse(31, "guide.md", content)

    assert result.file_type == "MARKDOWN"
    assert [block.section for block in result.blocks] == [
        "Guide",
        "Guide",
        "Guide > Example",
        "Guide > Example",
        "Guide > Example",
    ]
    code = result.blocks[3].text
    assert code == '```python\ndef hello():\n    return "ok"\n```'
    assert result.blocks[3].page_number is None


def test_markdown_setext_heading_is_recognized(parser: DocumentParser) -> None:
    result = parser.parse(32, "setext.markdown", b"Title\n=====\nBody\n")

    assert [block.text for block in result.blocks] == ["Title", "Body"]
    assert [block.section for block in result.blocks] == ["Title", "Title"]


def test_parses_utf8_text_without_inventing_structure(parser: DocumentParser) -> None:
    original = "第一行\n  indented second line"
    result = parser.parse(41, "notes.txt", original.encode("utf-8"))

    assert result.file_type == "TXT"
    assert result.blocks[0].text == original
    assert result.blocks[0].page_number is None
    assert result.blocks[0].section is None


def test_parses_utf8_bom(parser: DocumentParser) -> None:
    result = parser.parse(42, "bom.txt", b"\xef\xbb\xbfBOM text")

    assert result.blocks[0].text == "BOM text"


def test_rejects_invalid_text_encoding(parser: DocumentParser) -> None:
    with pytest.raises(TextDecodingError):
        parser.parse(43, "legacy.txt", b"\xff\xfe\x00\x80")


def test_rejects_unsupported_format(parser: DocumentParser) -> None:
    with pytest.raises(UnsupportedDocumentTypeError):
        parser.parse(51, "archive.zip", b"PK\x03\x04data")


def test_rejects_empty_and_oversized_files(parser: DocumentParser) -> None:
    with pytest.raises(EmptyDocumentError):
        parser.parse(52, "empty.txt", b"")

    small_limit_parser = DocumentParser(limits=ParsingLimits(max_file_size_bytes=3))
    with pytest.raises(DocumentTooLargeError):
        small_limit_parser.parse(53, "large.txt", b"four")


def test_parsed_block_is_not_a_final_chunk(parser: DocumentParser) -> None:
    result = parser.parse(54, "notes.txt", b"one document block")

    assert "chunk_index" not in result.blocks[0].model_dump()
